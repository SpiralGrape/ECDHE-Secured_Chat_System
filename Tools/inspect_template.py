import sys
from docx import Document
from docx.shared import Pt

path = r"C:\Users\陈浩然\Desktop\一些文件\大四上实践周\专业程序设计报告模板.docx"
d = Document(path)

print("=== sections ===")
for s in d.sections:
    print("page", s.page_width, s.page_height, "margins", s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)

print("\n=== styles (used) ===")
used = set()
for p in d.paragraphs:
    if p.text.strip():
        used.add(p.style.name)
print(sorted(used))

print("\n=== paragraphs ===")
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    if t:
        print(i, "|", p.style.name, "|", t[:80])

print("\n=== tables ===", len(d.tables))
for ti, tb in enumerate(d.tables):
    print("table", ti, "rows", len(tb.rows), "cols", len(tb.columns))
    for r in tb.rows[:4]:
        print("   ", [c.text.strip()[:20] for c in r.cells])

print("\n=== inline shapes ===", len(d.inline_shapes))
