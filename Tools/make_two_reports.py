# -*- coding: utf-8 -*-
"""按课程模板生成两份独立的课程设计报告（成员 A / 成员 B）。"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

TEMPLATE = r"C:\Users\陈浩然\Desktop\一些文件\大四上实践周\专业程序设计报告模板.docx"
OUTDIR = r"D:\ECDHE-Secured Chat System\docs"
IMG1 = os.path.join(os.environ.get("TEMP", "."), "mitm_demo", "client_1.png")
IMG2 = os.path.join(os.environ.get("TEMP", "."), "mitm_demo", "client_2.png")

doc = None


def new_doc():
    global doc
    doc = Document(TEMPLATE)
    body = doc.element.body
    for child in list(body):
        if child.tag == qn('w:sectPr'):
            continue
        body.remove(child)


# ------------------------------- 排版 helper -------------------------------
def _font(run, ascii_name="Times New Roman", ea_name="宋体", size=12, bold=None):
    run.font.name = ascii_name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), ea_name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def para(text, size=12, bold=False, align=None, indent=24, style='Normal',
         line=None, ea="宋体", ascii_name="Times New Roman", before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if indent:
        p.paragraph_format.first_line_indent = Pt(indent)
    if line:
        p.paragraph_format.line_spacing = line
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    _font(r, ascii_name, ea, size, bold)
    return p


def h1(text):
    return para(text, size=14, bold=True, indent=0, line=1.5, before=14, after=8)


def h2(text):
    return para(text, size=12, bold=True, indent=0, style='Body Text', before=8, after=4)


def body(text, indent=24):
    return para(text, size=12, indent=indent, line=1.5)


def code(lines):
    for ln in lines:
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(ln if ln else " ")
        _font(r, "Consolas", "宋体", 9)
    doc.add_paragraph()


def caption(text):
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    _font(r, "Times New Roman", "宋体", 10.5)
    return p


def table(headers, rows, size=10, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        _font(r, "Times New Roman", "宋体", size, True)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            r = p.add_run(str(v))
            _font(r, "Times New Roman", "宋体", size, False)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def picture(path, width_cm=13.5):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Cm(width_cm))
    else:
        para("（此处置入运行截图）", 10.5, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0)


# ------------------------------- 公共章节 -------------------------------
def cover(role):
    para("专业课程设计报告", size=36, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0,
         ea="黑体", ascii_name="黑体", before=40)
    para("（ 2026/ 2027 学年 第 一 学期）", size=15, align=WD_ALIGN_PARAGRAPH.CENTER,
         indent=0, ea="黑体", ascii_name="黑体", before=8)
    para("", size=12, indent=0)
    para("题   目：基于 ECDHE 的前向保密聊天工具", size=22, indent=54,
         ea="黑体", ascii_name="黑体", before=30)
    para("", size=12, indent=0)
    para("姓   名：________________", size=16, indent=54, ea="黑体", ascii_name="黑体", before=24)
    para("学   号：________________", size=16, indent=54, ea="黑体", ascii_name="黑体", before=8)
    para("专   业：________________", size=16, indent=54, ea="黑体", ascii_name="黑体", before=8)
    para("指导教师：________________", size=16, indent=54, ea="黑体", ascii_name="黑体", before=8)
    para("本人分工：" + role, size=14, indent=54, ea="黑体", ascii_name="黑体", before=8)
    doc.add_page_break()


def chapter_title():
    para("基于 ECDHE 的前向保密聊天工具", size=16, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, before=6, after=12)


def abstract(text, keywords):
    h2("摘  要")
    body(text)
    body("关键词：" + keywords, indent=0)


def ch1_requirement():
    h1("一、课题内容和要求")
    h2("课题内容")
    body("本课题围绕“前向保密”这一密码学安全属性展开。前向保密指长期主密钥泄露不会导致历史会话密钥泄露，"
         "即使系统遭受主动攻击也是如此。课题要求设计并实现具备前向安全性的即时通信工具：通信双方在通信前通过"
         "椭圆曲线迪菲-赫尔曼临时密钥交换（ECDHE）协商出本次会话专用密钥，并用该密钥对聊天内容做对称加密。"
         "程序要求使用 Windows 图形界面实现，支持多客户端同时在线并点对点通信，中继服务器不得具备解密能力。")
    h2("课题要求")
    body("（1）理解消息安全通信工具的基本概念与工作原理，明确机密性、完整性、身份认证与前向保密之间的联系与区别。")
    body("（2）理解前向保密的基础理论，掌握 ECDHE 能够提供前向保密的原因，以及与 RSA 密钥交换的本质差异。")
    body("（3）采用传输层安全协议（TLS）中的相关密码套件，整合 ECDHE 等支持前向安全性的算法，保证每次会话生成独立会话密钥。")
    body("（4）完成可运行、可演示的图形化聊天工具，并通过实验证明：长期主密钥泄露后，历史通信依然安全。")


def ch2_requirement_analysis():
    h1("二、需求分析")
    h2("威胁模型")
    body("在设计之前先明确系统面对的威胁，再据此推导需求。系统需要应对的攻击者能力及对应设计如下表。")
    table(["攻击者能力", "系统设计响应"], [
        ["被动窃听链路（抓包）", "全程 AEAD 加密，攻击者只能得到密文与长度"],
        ["攻破中继服务器", "服务器不持有任何密钥，无法解密；篡改会破坏 AAD 认证"],
        ["主动中间人（替换公钥）", "Ed25519 签名 + Finished 校验阻断篡改；首次接触靠安全码核对"],
        ["重放、篡改、乱序", "帧头（会话号、序号、收发方）作为 AAD 参与认证，序号严格单调"],
        ["事后获取长期主密钥", "会话密钥只由临时私钥派生，历史会话不可还原（前向保密）"],
        ["当前棘轮密钥被窃取", "每 8 条消息触发 DH 棘轮轮换，旧私钥即时销毁（后向恢复）"],
        ["本机文件被拷贝", "身份私钥 DPAPI 加密，聊天记录 AES-GCM 加密，换账户不可解密"],
        ["进程崩溃、断网", "心跳 15 秒判定超时并自动重连，服务器 60 秒回收僵尸连接"],
    ], size=9.5, widths=[5.5, 10])

    h2("功能需求")
    body("① 多客户端在线，按 ID 点对点通信；② 消息端到端加密且服务器不可解密；③ 会话列表与聊天记录持久化，"
         "重启后自动重载；④ 断线自动重连、离线提示与发送失败反馈；⑤ 可验证的对端身份（安全码与指纹固定）。")

    h2("非功能需求")
    body("安全性：机密性、完整性、前向保密、抗重放，以及配合安全码的抗中间人能力。"
         "性能：单次握手在局域网内小于 10 毫秒。可用性：断网后 15 秒内发现并自动恢复。"
         "可维护性：密码学层与网络层解耦，协议实现可被测试工具复用。")


def ch3_overall():
    h1("三、系统总体设计")
    h2("总体架构")
    body("系统采用客户端—中继服务器—客户端的结构。客户端承担全部密码学运算；中继服务器只按目标 ID 转发密文帧，"
         "不持有任何密钥，属于“不可信中继”；因此服务器被攻破也不会泄露聊天内容。")
    table(["层次", "模块", "职责"], [
        ["密码学层", "crypto_utils.h", "X25519、HKDF-SHA256、AES-256-GCM、Ed25519、密钥清零"],
        ["协议层", "protocol.h / secure_session.h", "帧格式、AAD 构造、握手、消息级棘轮、DH 棘轮、安全码"],
        ["中继层", "server.cpp", "连接管理、按 ID 路由、心跳应答、超时回收"],
        ["客户端", "ClientGUI.cpp", "会话管理、界面、加密持久化、自动重连"],
        ["测试与工具", "Tests/ Tools/", "28 项自测、端到端脚本、中间人攻击工具"],
    ], size=9.5, widths=[3, 5, 8])
    picture(IMG1, 12.5)
    caption("图 1 客户端主界面与安全状态栏")

    h2("通信帧格式")
    body("TCP 上每条消息统一为“长度(4) + 目标/源 ID(4) + 载荷”，载荷前两字节为版本与类型。主要类型如下表。")
    table(["类型", "值", "长度（字节）", "说明"], [
        ["HS_INIT", "0x01", "162", "握手请求（临时公钥、随机数、身份公钥、签名）"],
        ["HS_RESP", "0x02", "186", "握手响应（含签名与 MAC_B）"],
        ["HS_CONFIRM", "0x03", "42", "握手确认（MAC_A）"],
        ["REKEY", "0x04", "66", "DH 棘轮轮换（加密传输，载荷为新公钥）"],
        ["DATA", "0x10", "18+密文+16", "加密聊天数据"],
        ["PING/PONG", "0x20/0x21", "2", "心跳（目标 ID 为 0，不转发）"],
        ["SYSTEM", "0xFF", "6+文本", "服务器提示（如用户不在线）"],
    ], size=9.5, widths=[3, 2.5, 3, 8.5])

    h2("数据流与本地存储")
    body("发送路径：界面输入 → 会话层加密（Seal）→ 组帧 → TCP → 中继转发 → 对端解密（Open）→ 界面显示。"
         "接收路径反之。聊天记录按“本地身份指纹/对端 ID”分文件保存，每条记录用身份私钥派生的密钥做 AES-GCM 加密；"
         "身份私钥本身用 Windows DPAPI 加密存储，保证换账户或换机器都无法解密。")


# ------------------------------- 测试章节（公共） -------------------------------
def ch7_test():
    h1("七、系统测试与结果分析")
    h2("测试环境")
    table(["项目", "配置"], [
        ["操作系统", "Windows 11 (x64)"],
        ["编译器 / 构建", "MSVC 19.3x（VS2022）、CMake 3.20+、vcpkg（openssl 3.x、boost-asio）"],
        ["网络", "本机回环 127.0.0.1:9000"],
        ["测试程序", "CryptoTest.exe、Tools\\smoke_test.ps1、Tools\\mitm_demo.ps1"],
    ], size=9.5, widths=[4, 11])

    h2("功能测试用例")
    table(["编号", "测试项", "输入 / 操作", "预期结果", "实际结果"], [
        ["F-01", "启动与连接", "启动服务器后启动客户端", "标题显示用户 ID，状态栏显示已连接", "一致"],
        ["F-02", "端到端加密消息", "A 向 2 发送消息", "B 正确显示发送方与内容", "一致"],
        ["F-03", "双向通信", "B 回复 A", "A 正确显示", "一致"],
        ["F-04", "会话建立提示", "首次发送", "状态栏显示安全码与密钥指纹", "一致"],
        ["F-05", "历史持久化", "关闭后重新打开客户端", "自动加载历史并选中最近会话", "一致"],
        ["F-06", "落盘为密文", "查看 chats\\*.log", "内容为十六进制密文，无明文", "一致"],
        ["F-07", "解密导出", "--dump 导出", "导出文件包含完整明文", "一致"],
        ["F-08", "多档案隔离", "--profile p1 / p2", "身份与历史互不干扰", "一致"],
        ["F-09", "离线用户提示", "向不在线的 ID 发送", "状态栏红色提示，聊天区出现警示", "一致"],
        ["F-10", "发给自己", "收件人为本人", "提示不能给自己发消息", "一致"],
        ["F-11", "空收件人", "收件人留空", "提示请先填写收件人 ID", "一致"],
        ["F-12", "超长消息", "发送 70000 字符", "提示超过 64 KB 并取消发送，连接不断", "一致"],
        ["F-13", "会话列表与未读", "收到非当前会话的消息", "列表出现未读标记并支持切换", "一致"],
        ["F-14", "窗口自适应", "拉伸窗口", "控件重新布局且无重叠", "一致"],
    ], size=8.5, widths=[1.6, 3, 4.5, 5.5, 1.8])

    h2("健壮性测试")
    table(["编号", "测试项", "方法", "预期结果", "实际结果"], [
        ["R-01", "服务器中途关闭", "运行中结束服务器进程", "15 秒内提示断开并进入重试", "一致"],
        ["R-02", "服务器重启", "重新启动服务器", "自动重连并恢复可用", "一致"],
        ["R-03", "客户端异常退出", "强制结束客户端", "服务器 60 秒内回收该 ID", "已实现"],
        ["R-04", "畸形帧", "发送长度 0xFFFFFFFF 的帧", "服务器断开该连接，不影响他人", "一致"],
        ["R-05", "篡改密文", "翻转密文 1 bit", "认证失败并重置会话", "一致"],
        ["R-06", "篡改帧头", "修改序号或收发方", "解密失败", "一致"],
        ["R-07", "并发会话", "三个客户端两两通信", "消息不串会话", "设计保证"],
        ["R-08", "同时发起握手", "双方同时发送", "按 ID 小者胜出，仅一条会话", "设计保证"],
        ["R-09", "棘轮压力", "连续发送 10 条以上", "触发 REKEY 后仍按序送达", "一致"],
    ], size=8.5, widths=[1.6, 3, 4.5, 5.5, 1.8])

    h2("安全专项测试")
    table(["编号", "测试项", "方法", "预期结果", "实际结果"], [
        ["S-01", "中间人攻击", "运行 mitm_demo 脚本", "中间人打印明文，两端安全码不同", "一致"],
        ["S-02", "前向保密", "记录密文后删除 identity.pem 重启", "历史密文无法解密", "一致"],
        ["S-03", "后向恢复", "连发超过 8 条触发 DH 轮换", "旧链密钥无法解密新消息", "一致"],
        ["S-04", "重放攻击", "重复投递同一 DATA 帧", "第二次被序号校验拒绝", "一致"],
        ["S-05", "指纹变化阻断", "先经中间人建立信任再连真服务器", "指纹不一致被阻断", "已实现"],
        ["S-06", "服务器不可解密", "查看服务器日志", "只有字节数，无明文", "一致"],
    ], size=8.5, widths=[1.6, 3, 4.5, 5.5, 1.8])

    h2("中间人攻击实验结果")
    body("中间人工具冒充中继服务器，对每一对通信者分别与两端各完成一次握手，因此同时握有两把会话密钥。实测输出：")
    code([
        "[MITM] 已冒充用户 2 接受用户 1 的握手",
        "[MITM] 已冒充用户 1 向用户 2 发起握手",
        "[MITM] 欺骗成功：用户 1 与用户 2 都显示“加密会话已建立”",
        "[MITM] 但两端的安全码不同：用户 1 看到 629420，用户 2 看到 795209",
        "[截获] 1 -> 2 明文: card-password-is-123456",
        "[抓包]   密文 39 字节，前 16 字节 fdb3776657d7af05ca5e9bc670fd210e",
    ])
    picture(IMG1, 11.5)
    caption("图 2 一端客户端显示的安全码与会话状态")
    picture(IMG2, 11.5)
    caption("图 3 另一端客户端显示的安全码（与图 2 不同）")
    body("结论：首次接触且中间人在场时，ECDHE 无法自动阻止劫持，这是信息论限制而非实现缺陷；"
         "但两端安全码必然不同，核对一次即可发现。这与 Signal、WhatsApp 的安全模型一致。")

    h2("测试结论")
    body("（1）28 项密码学与协议断言全部通过，覆盖签名、密钥交换、KDF、AEAD、两种棘轮与完整握手流程。")
    body("（2）端到端自动化测试覆盖“握手 → 收发 → 落盘 → 导出 → 棘轮”全链路，5 项断言全部通过。")
    body("（3）边界与异常路径（离线、自发、空收件人、超长消息、断线重连、畸形帧）均给出明确提示，未出现崩溃或卡死。")
    body("（4）安全专项验证了服务器不可解密、前向保密、后向恢复、重放拒绝与中间人可被安全码发现五项结论。")
    body("（5）已知不足：不支持乱序与丢帧恢复，弱网丢包会导致会话重置（见第八章局限性）。")


# ------------------------------- 成员 A -------------------------------
def personal_a():
    h1("四、ECDHE 握手协议与密钥调度设计")
    body("本章由成员 A 负责，说明握手的三个报文如何协商出会话密钥，以及签名与密钥派生为什么这样设计。")

    h2("三报文握手流程")
    body("发起方 A 生成临时 X25519 密钥对与随机数 nonceA，发送 INIT；响应方 B 验证签名后生成自己的临时密钥对与 nonceB，"
         "计算共享秘密 Z 并派生链密钥，回复 RESP；A 验证 RESP 后发送 CONFIRM，双方会话建立。"
         "整个过程服务器只做原样转发。")
    body("签名的覆盖范围是防篡改的关键：Sig_A 覆盖版本、类型、会话号、pubA、nonceA、身份公钥以及双方 ID；"
         "Sig_B 覆盖 RESP 前 90 字节再加上 Sig_A，从而把 B 的响应绑定到 A 的本次请求，防止报文被拼接或跨会话重放。")

    h2("HKDF 密钥派生与域分离")
    body("共享秘密先经 HKDF-Extract 条件化，再按不同 info 标签分别派生出“发起方到响应方”链密钥、"
         "“响应方到发起方”链密钥与 Finished 密钥。域分离保证收发两个方向使用不同的密钥，"
         "同时把 Finished 密钥与消息密钥彻底隔离。")
    code([
        "Z        = X25519(临时私钥, 对端临时公钥)",
        "salt     = nonceA || nonceB",
        "secret   = HKDF-SHA256(Z, salt, \"CHAT-v1 ECDHE X25519\")",
        "chainI2R = HKDF(secret, \"\", \"chain initiator->responder\")",
        "chainR2I = HKDF(secret, \"\", \"chain responder->initiator\")",
        "finished = HKDF(secret, \"\", \"finished key\")",
        "root     = HKDF(secret, \"\", \"root key\")      // DH 棘轮起点",
    ])

    h2("Finished 校验与安全码")
    body("MAC 的作用不是身份认证（那由签名完成），而是密钥确认：双方用 finished_key 对 transcript 做 HMAC，"
         "证明派生出完全相同的会话密钥。握手完成后再由 transcript 派生 6 位安全码，显示在状态栏供双方口头核对。")

    h1("五、密钥棘轮与安全属性")
    h2("消息级密钥棘轮")
    body("每发送一条消息，先由链密钥派生消息密钥与 nonce，再用 HKDF 推进链密钥并覆盖旧值：")
    code([
        "msg_key = HKDF(chain, \"\", \"message key\")",
        "nonce   = HKDF(chain, \"\", \"message iv\")",
        "chain'  = HKDF(chain, \"\", \"chain next\")     // 立即覆盖旧 chain",
    ])
    body("由此获得消息级前向保密：即使当前链密钥泄露，也无法还原此前消息的密钥。")

    h2("DH 棘轮（后向恢复）")
    body("每发送 8 条消息，发送方生成新的 X25519 密钥对，用旧链加密一条 REKEY 控制帧携带新公钥；"
         "双方随后以 root' = HKDF(salt = root, ikm = DH(新私钥, 对端当前公钥)) 更新根密钥并重建双向链，"
         "发送方旧私钥被覆盖销毁。因此攻击者即使窃取了某一时刻的棘轮私钥，受害方下一次轮换后也无法继续推导新密钥。")
    code([
        "if (send_since_rekey_ >= REKEY_INTERVAL) {",
        "    auto fresh = x25519_generate();",
        "    Bytes freshPub = x25519_public_raw(fresh.get());",
        "    BuildRecord(proto::REKEY, freshPub, myId, peerId, rec, err);  // 旧链加密",
        "    out.push_back(rec);",
        "    Bytes z = x25519_shared(fresh.get(), peer.get());",
        "    ApplyRatchet(z, err);            // 更新根密钥并重建双向链",
        "    dh_self_ = std::move(fresh);     // 旧私钥被销毁",
        "    send_since_rekey_ = 0;",
        "}",
    ])

    h2("密钥生命周期与清零")
    body("所有中间值（Z、secret、链密钥、消息密钥）在派生出下一层后立即调用 OPENSSL_cleanse 清零；"
         "临时私钥在首次 DH 轮换后被销毁。需要注意的是“清零”与“使用”的先后顺序——本项目曾出现过根密钥"
         "在缓冲区清零之后才派生的问题，功能正常但强度归零，说明这类缺陷只能靠代码审查与断言发现。")

    h1("六、安全性分析")
    h2("威胁逐条对照")
    table(["威胁", "是否可防", "依据"], [
        ["链路窃听", "可防", "AEAD 加密，只泄露长度与收发关系"],
        ["服务器解密", "可防", "服务器无密钥（中间人实验反证：只有主动劫持才能解密）"],
        ["篡改帧头", "可防", "帧头作为 AAD 参与认证"],
        ["重放", "可防", "序号严格单调，失步即作废会话"],
        ["事后主密钥泄露", "可防", "会话级与消息级前向保密"],
        ["当前密钥被窃取", "可防（一次轮换后）", "DH 棘轮轮换并销毁旧私钥"],
        ["首次接触的中间人", "可发现，不能自动阻止", "签名保证未被篡改，但身份需安全码核对"],
    ], size=9.5, widths=[4, 4, 7])

    h2("与 TLS 1.3 / Signal 的对比")
    table(["特性", "本系统", "TLS 1.3", "Signal"], [
        ["密钥交换", "X25519 ECDHE", "X25519 / P-256", "X3DH"],
        ["身份认证", "Ed25519 + 安全码", "证书链（PKI）", "身份密钥 + 安全码"],
        ["前向保密", "会话级 + 消息级", "会话级", "会话级 + 消息级"],
        ["后向恢复", "周期性 DH 棘轮", "KeyUpdate", "DH 棘轮"],
        ["乱序容忍", "不支持", "支持（流式）", "支持（缓存跳跃密钥）"],
    ], size=9.5, widths=[3, 5, 3.5, 3.5])

    h2("本人负责的测试")
    body("（1）28 项密码学与协议断言，覆盖签名、X25519、HKDF、AES-GCM、消息级棘轮，以及“完整握手 + 20 条消息 + "
         "2 次自动轮换”的协议级用例。")
    body("（2）安全专项中的篡改、重放、前向保密与后向恢复判定逻辑，均由本人设计断言并在自测中验证。")
    body("（3）中间人实验的结论解释：两端安全码由各自的 transcript 派生，中间人插入两段独立握手后必然不同。")


# ------------------------------- 成员 B -------------------------------
def personal_b():
    h1("四、网络协议与并发设计")
    body("本章由成员 B 负责，说明帧如何在中继与客户端之间可靠传递，以及并发与异常路径如何处理。")

    h2("帧格式与错误处理")
    body("服务器对每一帧先读取 4 字节长度并做上下界检查：小于 4 或超过 1 MiB 立即断开该连接，"
         "避免攻击者用一个长度字段耗尽内存。随后读取载荷，前 4 字节为目标 ID，其余为端到端加密的载荷，服务器不做任何解析。")

    h2("中继路由与每连接写锁")
    body("转发的难点在并发：A 的消息由“处理 A 的线程”写入 B 的 socket，而 B 自己的线程也可能同时写同一个 socket，"
         "若不互斥就会出现帧交错、对端无法还原帧边界。为此每个连接维护一把写锁，长度与载荷在同一临界区内连续写出。")
    body("同时，全局在线表只在锁内做查找，阻塞写放到锁外，避免一个慢客户端拖垮整个服务。")
    code([
        "if (len < 4 || len > proto::MAX_PAYLOAD) break;          // 防 OOM",
        "SessionPtr dst; { lock(g_clients_mu); dst = find(target); }   // 锁内只查找",
        "if (dst) {",
        "    std::lock_guard<std::mutex> lk(dst->write_mu);       // 每连接一把写锁",
        "    write(dst->sock, &len, 4); write(dst->sock, payload, len);",
        "}",
    ])

    h2("心跳与超时回收")
    body("客户端每 5 秒发送 PING（目标 ID 为 0，服务器本地应答 PONG，不参与转发），15 秒未收到 PONG 即判定断线并"
         "主动关闭 socket 触发重连。服务器每 10 秒巡检一次，对超过 60 秒无活动的连接主动断开，避免客户端崩溃后"
         "ID 被永久占用。")

    h1("五、客户端工程实现")
    h2("会话管理的调用路径")
    body("会话状态机与握手调用由成员 A 实现的 SecureSession 提供，本人在客户端负责：维护每个对端的会话对象与待发队列、"
         "在未建立会话时触发握手并把消息排队、会话建立后按序补发。该部分与网络线程通过互斥量协作，"
         "锁序固定为“会话锁 → 网络锁”，避免死锁。")

    h2("加密持久化与 DPAPI")
    body("聊天记录按“本地身份指纹 / 对端 ID”分文件保存，每条记录用身份私钥经 HKDF 派生的密钥做 AES-256-GCM 加密，"
         "并把对端 ID 作为 AAD；直接查看文件只能看到十六进制密文。身份私钥本身用 Windows DPAPI 加密，"
         "文件被拷贝到其他账户或机器都无法解密。")
    body("归档键选择“身份指纹”而不是“服务器分配的 ID”，是因为 ID 是临时的（重启即变），身份密钥才是持久的；"
         "这样重启或换 ID 后历史与指纹固定仍然有效。")

    h2("多档案与导出")
    body("通过 --profile 参数让同一台机器上的多个实例拥有独立的身份、历史与指纹目录；通过 --dump 参数在本地解密并"
         "导出聊天记录，既可用于备份，也用于自动化测试的断言。")

    h2("界面与状态栏")
    body("界面分三块：左侧会话列表（未读标记）、中间聊天区、底部输入行与状态栏。握手、指纹、断线等信息只写入状态栏，"
         "聊天区只保留消息，避免刷屏；发送给离线用户时聊天区出现警示行、状态栏变红。")

    h1("六、测试工程与工具链")
    h2("端到端自动化")
    body("由于跨进程读取多行编辑框不可靠，测试脚本改为通过窗口标题识别客户端 ID、按控件样式位定位控件，"
         "驱动界面发送消息后，用 --dump 导出解密后的历史做断言。这一设计避免了“读不到控件就误判为空”的假成功。")

    h2("中间人工具的设计")
    body("攻击工具复用与客户端完全相同的 SecureSession 实现，冒充服务器：收到 A 的 INIT 后，一方面以 B 的身份"
         "回应 A，另一方面以 A 的身份向 B 发起握手，于是同时握有两把会话密钥，可解密、打印并重新加密转发。"
         "握手未完成时到达的数据帧会进入待转发队列，待双方就绪后按序补发。")

    h2("边界与健壮性测试")
    body("针对离线用户、发给自己、空收件人、超长消息、服务器中途关闭、畸形帧等场景逐一验证，确保都有明确提示且不崩溃。"
         "其中离线提示曾因服务器系统消息帧缺少 4 字节发送者前缀而完全失效，是通过“客户端加日志发现分支从未命中”定位的。")

    h2("构建与依赖")
    body("使用 CMake + vcpkg 管理依赖。CMake 4 已移除 FindBoost 模块，必须改用 Boost 的 CONFIG 模式；"
         "源码统一以 UTF-8 编译并配合 W 版 API，解决了中文乱码；客户端链接 crypt32 以使用 DPAPI。")


# ------------------------------- 第八章 -------------------------------
def ch8_summary(personal_para):
    h1("八、总结与展望")
    body("本次课程设计围绕“基于 ECDHE 的前向保密聊天工具”展开。通过查阅资料，明确了前向保密的本质是"
         "把长期密钥与一次性会话密钥解耦，从而在长期密钥泄露时保护历史通信；也理解了 ECDHE 只是实现手段，"
         "前向保密才是需要证明的安全属性。")
    body("系统采用 C++17 与 OpenSSL 3.x，自行实现了 TLS 1.3 风格的握手：X25519 临时密钥交换、HKDF-SHA256 域分离派生、"
         "AES-256-GCM 记录层与 Ed25519 身份签名，并叠加消息级密钥棘轮与周期性 DH 棘轮，"
         "同时获得前向保密与后向恢复能力。中继服务器被设计为不可信中继，任何时刻都不持有密钥。")
    body(personal_para)
    body("系统仍有明确的边界：不支持乱序与丢帧恢复，身份信任依赖人工核对，元数据对服务器可见，"
         "未实现离线消息、群聊与文件传输。后续可通过每条消息携带 DH 头的完整双棘轮、KeyUpdate、离线消息队列、"
         "消息填充与形式化验证继续完善。")

    h2("本人总结")
    body("通过本次设计，我不仅掌握了所用算法的具体用法，更重要的是建立了“先建模威胁、再选择机制、"
         "最后用实验证伪”的安全工程思维：任何一个“看起来能用”的实现，都必须用可失败的测试去证明。")


# ------------------------------- 生成 -------------------------------
def build(role, abstract_text, keywords, personal, personal_para, filename):
    new_doc()
    cover(role)
    chapter_title()
    abstract(abstract_text, keywords)
    ch1_requirement()
    ch2_requirement_analysis()
    ch3_overall()
    personal()
    ch7_test()
    ch8_summary(personal_para)
    out = os.path.join(OUTDIR, filename)
    doc.save(out)
    print("saved:", out, os.path.getsize(out))


def main():
    os.makedirs(OUTDIR, exist_ok=True)

    build(
        "密码协议设计（成员 A）",
        "本报告实现了一个具备前向保密与后向恢复能力的端到端加密聊天工具。系统以 X25519 的 ECDHE 临时密钥交换为核心，"
        "每次会话都协商出独立的会话密钥，并通过 HKDF-SHA256 按用途做域分离；记录层采用 AES-256-GCM，"
        "把帧头作为附加认证数据参与完整性校验；在此之上叠加消息级密钥棘轮与周期性 DH 棘轮，"
        "使历史消息不可被事后解密、被泄露的当前密钥也无法推导未来密钥。身份层使用 Ed25519 签名与 6 位安全码，"
        "并实现了一个可运行的中间人攻击工具来定量验证安全边界。",
        "前向保密；ECDHE；X25519；HKDF；AES-GCM；Ed25519；密钥棘轮；端到端加密",
        personal_a,
        "本人主要负责密码协议部分：设计并实现了三报文握手、签名覆盖范围、HKDF 密钥调度、Finished 校验与安全码派生，"
        "以及消息级棘轮与 DH 棘轮两项密钥更新机制，并编写了 28 项密码学与协议断言。"
        "在这个过程中，最深的体会是密码学实现的正确性往往取决于“细节的对称性”——签名与验签的输入、"
        "派生顺序、清零时机，任何一处不一致都会导致握手失败或安全性归零，而且往往不会报错。",
        "专业程序设计报告-成员A-密码协议方向.docx")

    build(
        "网络与工程化（成员 B）",
        "本报告实现了一个具备前向保密能力的端到端加密聊天工具。系统采用 C/S 结构：客户端承担全部密码学运算，"
        "中继服务器只按目标 ID 转发密文帧、应答心跳并回收超时连接，不持有任何密钥。"
        "工程上重点解决了并发写同一 socket 的帧交错、锁序与全局锁阻塞、心跳与断线重连、"
        "聊天记录的加密持久化与身份私钥的 DPAPI 保护，并配套实现了端到端自动化测试与中间人攻击演示工具。",
        "端到端加密；ECDHE；AES-GCM；并发；心跳重连；加密存储；DPAPI；自动化测试",
        personal_b,
        "本人主要负责网络与工程化部分：中继服务器的路由与并发控制、客户端的心跳与自动重连、加密持久化与 DPAPI、"
        "多档案与导出功能，以及端到端自动化测试和中间人演示工具。"
        "这部分的工作让我体会到：安全系统的可靠性往往不取决于算法，而取决于异常路径——断线、丢帧、畸形包、"
        "中文路径、进程占用，这些“非功能”细节才是最容易出问题的地方。",
        "专业程序设计报告-成员B-网络与工程化方向.docx")


if __name__ == "__main__":
    main()
